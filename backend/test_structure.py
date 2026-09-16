"""G1 regressions. Language preservation is not multilingual legal accuracy."""
import copy
import hashlib
from io import BytesIO
from pathlib import Path
import pytest
from pydantic import ValidationError
from backend.domain import ContractReport, validate_contract_report
from backend.engine import analyze, extract, DEMO

FIXTURE = Path(__file__).parent / "fixtures" / "multilingual_structure.txt"

def report(text=DEMO):
    return analyze(extract(text.encode("utf-8"),"fixture.txt"),"fixture.txt","2026-09-13")

def test_version_hash_and_legacy_shape():
    r=report()
    assert r["schema_version"]=="1.0"
    assert r["document"]["sha256"]==hashlib.sha256(DEMO.encode()).hexdigest()
    assert r["document"]["hash_basis"]=="original_bytes"
    assert r["document"]["parser_version"]=="structure-0.2.0"
    assert r["rule_version"]=="patterns-0.1.0"
    assert r["findings"][0]["confidence"] is None
    assert "semantic" not in r["findings"][0]
    assert validate_contract_report(r)==r

def test_unicode_hierarchy_and_scoped_labels():
    r=report(FIXTURE.read_text(encoding="utf-8"))
    body={c["label"]:c for c in r["clauses"] if c["scope"]=="body"}
    assert body["1.1(a)"]["parent_id"]==body["1.1"]["id"]
    assert body["7(b)(ii)"]["parent_id"]==body["7(b)"]["id"]
    assert body["3.1"]["parent_id"]==body["3"]["id"]
    assert "வாங்குபவர்" in body["2.1"]["text"]
    assert "खरीदार" in body["3.1"]["text"]
    assert r["document"]["scripts"]==["Devanagari","Latin","Tamil"]
    assert len([c for c in r["clauses"] if c["label"]=="1.1"])==2
    assert any("primarily English" in s for s in r["limitations"])

def test_every_clause_reconstructs_from_source():
    r=report(FIXTURE.read_text(encoding="utf-8"))
    units={u["id"]:u for u in r["document"]["units"]}
    for c in r["clauses"]:
        assert c["text"]=="\n".join(units[s["unit_id"]]["text"][s["start"]:s["end"]] for s in c["spans"])
    assert r["unassigned_spans"]
    assert r["document"]["coverage"]=="partial"

@pytest.mark.parametrize("newline",["\n","\r\n","\r","\u2028"])
def test_offsets_and_whitespace(newline):
    r=report(newline.join(["  Preamble","  1 Heading  ","    Body தமிழ் हिन्दी  ","  2 Next","Text."]))
    assert r["clauses"][0]["text"]=="1 Heading\nBody தமிழ் हिन्दी"
    assert validate_contract_report(r)

def test_date_line_not_a_heading():
    r=report("1 Dates\n1 January 2027 is the start date.\n2 Other\nSome terms.")
    assert len(r["clauses"])==2
    assert "1 January 2027" in r["clauses"][0]["text"]

def test_ambiguous_relative_numbering_remains_visible():
    r=report("1 Terms\n(a) General\n(i) Ambiguous nesting\nSome obligation.")
    assert any("Unresolved relative numbering" in s for s in r["limitations"])
    assert "(i) Ambiguous" in r["clauses"][-1]["text"]

def test_duplicate_numbering_does_not_guess():
    r=report("1 One\nA.\n1 Duplicate\nB.\n2 Uses\nSee Section 1.")
    assert r["references"][0]["status"]=="UNRESOLVED"
    assert any("Duplicate clause label" in s for s in r["limitations"])

def test_docx_table_locations_and_multilingual_text():
    from docx import Document
    doc=Document();doc.add_paragraph("1 Terms")
    table=doc.add_table(rows=1,cols=2)
    table.cell(0,0).text="தமிழ்";table.cell(0,1).text="हिन्दी"
    stream=BytesIO();doc.save(stream)
    r=analyze(extract(stream.getvalue(),"fixture.docx"),"fixture.docx")
    cells=[u for u in r["document"]["units"] if u["kind"]=="table_cell"]
    assert [c["location"]["cell"] for c in cells]==[1,2]
    assert all(c["page"] is None for c in r["clauses"])
    assert "தமிழ்" in r["clauses"][0]["text"]
    assert "हिन्दी" in r["clauses"][0]["text"]

def test_docx_automatic_numbering_warns():
    from docx import Document
    doc=Document();doc.add_paragraph("1 Heading");doc.add_paragraph("Auto list item",style="List Number")
    stream=BytesIO();doc.save(stream)
    r=analyze(extract(stream.getvalue(),"numbered.docx"),"numbered.docx")
    assert any("automatic list numbering" in w for w in r["document"]["warnings"])

def test_blank_pdf_page_is_not_a_scan_failure():
    import pymupdf
    doc=pymupdf.open();p=doc.new_page();p.insert_text((50,50),"1 Terms\nThis is a readable source clause.")
    doc.new_page()
    r=analyze(extract(doc.tobytes(),"blank-page.pdf"),"blank-page.pdf")
    assert len(r["document"]["units"])==2
    assert r["clauses"][0]["page_end"]==1
    doc.close()

@pytest.mark.parametrize("mutation",["quote","offset","line","document","parent","finding","summary","omission","overlap","version"])
def test_corrupt_reports_rejected(mutation):
    r=report()
    if mutation=="quote":r["clauses"][0]["text"]+=" invented"
    elif mutation=="offset":r["clauses"][0]["spans"][0]["end"]=99999999
    elif mutation=="line":r["clauses"][0]["spans"][0]["line"]=999
    elif mutation=="document":r["clauses"][0]["document_id"]="other"
    elif mutation=="parent":r["clauses"][0]["parent_id"]=r["clauses"][0]["id"]
    elif mutation=="finding":r["findings"][0]["evidence"]=["nonexistent"]
    elif mutation=="summary":r["summary"]["clauses"]=999
    elif mutation=="omission":r["unassigned_spans"]=[]
    elif mutation=="overlap":r["unassigned_spans"].append(copy.deepcopy(r["clauses"][0]["spans"][0]))
    elif mutation=="version":r["schema_version"]="99"
    with pytest.raises(ValidationError):ContractReport.model_validate(r)

def test_text_limit_fails_before_analysis():
    with pytest.raises(ValueError,match="limit"):
        extract(("1 Heading\n"+"a"*1_000_001).encode(),"large.txt")
